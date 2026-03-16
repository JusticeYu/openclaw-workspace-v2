#!/usr/bin/env python3
"""
生成季度立项规划报告PPT模板
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


def create_planning_template(output_path: str = "planning_template.pptx"):
    """创建规划报告PPT模板"""
    
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # 定义颜色
    PRIMARY_COLOR = RGBColor(0, 82, 155)  # 深蓝色
    ACCENT_COLOR = RGBColor(0, 150, 214)  # 浅蓝色
    TEXT_COLOR = RGBColor(51, 51, 51)  # 深灰色
    
    # ===== 第1页：封面 =====
    slide_layout = prs.slide_layouts[6]  # 空白布局
    slide = prs.slides.add_slide(slide_layout)
    
    # 标题
    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(11.333), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.text = "【项目名称】季度立项规划报告"
    title_frame.paragraphs[0].font.size = Pt(44)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = PRIMARY_COLOR
    title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 副标题
    subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.333), Inches(1))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "202X年第X季度"
    subtitle_frame.paragraphs[0].font.size = Pt(28)
    subtitle_frame.paragraphs[0].font.color.rgb = ACCENT_COLOR
    subtitle_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 信息
    info_box = slide.shapes.add_textbox(Inches(1), Inches(5.5), Inches(11.333), Inches(1))
    info_frame = info_box.text_frame
    info_frame.text = "PO：【姓名】\n日期：202X年XX月XX日"
    info_frame.paragraphs[0].font.size = Pt(18)
    info_frame.paragraphs[0].font.color.rgb = TEXT_COLOR
    info_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # ===== 第2页：目录 =====
    slide = prs.slides.add_slide(slide_layout)
    
    # 页面标题
    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    header_frame = header_box.text_frame
    header_frame.text = "目录"
    header_frame.paragraphs[0].font.size = Pt(32)
    header_frame.paragraphs[0].font.bold = True
    header_frame.paragraphs[0].font.color.rgb = PRIMARY_COLOR
    
    # 目录内容
    toc_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(11), Inches(5))
    toc_frame = toc_box.text_frame
    toc_items = [
        "1. 季度目标与策略",
        "2. 特性规划",
        "3. 资源与负荷分析",
        "4. 里程碑与关键节点",
        "5. 风险与应对措施"
    ]
    for i, item in enumerate(toc_items):
        if i == 0:
            toc_frame.text = item
        else:
            p = toc_frame.add_paragraph()
            p.text = item
            p.font.size = Pt(24)
            p.font.color.rgb = TEXT_COLOR
            p.space_before = Pt(20)
    
    toc_frame.paragraphs[0].font.size = Pt(24)
    toc_frame.paragraphs[0].font.color.rgb = TEXT_COLOR
    
    # ===== 第3页：季度目标 =====
    slide = prs.slides.add_slide(slide_layout)
    
    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    header_frame = header_box.text_frame
    header_frame.text = "1. 季度目标与策略"
    header_frame.paragraphs[0].font.size = Pt(32)
    header_frame.paragraphs[0].font.bold = True
    header_frame.paragraphs[0].font.color.rgb = PRIMARY_COLOR
    
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(5.5))
    content_frame = content_box.text_frame
    content_frame.text = "季度目标："
    
    p = content_frame.add_paragraph()
    p.text = "• 【在此填写季度主要业务目标】"
    p.font.size = Pt(20)
    p.space_before = Pt(12)
    
    p = content_frame.add_paragraph()
    p.text = "• 【例如：完成XX功能上线，提升用户体验】"
    p.font.size = Pt(20)
    p.space_before = Pt(8)
    
    p = content_frame.add_paragraph()
    p.text = "\n策略重点："
    p.font.size = Pt(22)
    p.font.bold = True
    p.space_before = Pt(20)
    
    p = content_frame.add_paragraph()
    p.text = "• 【策略1：例如聚焦核心业务功能】"
    p.font.size = Pt(20)
    p.space_before = Pt(12)
    
    p = content_frame.add_paragraph()
    p.text = "• 【策略2：例如技术债务清理】"
    p.font.size = Pt(20)
    p.space_before = Pt(8)
    
    content_frame.paragraphs[0].font.size = Pt(22)
    content_frame.paragraphs[0].font.bold = True
    
    # ===== 第4页：特性规划 =====
    slide = prs.slides.add_slide(slide_layout)
    
    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    header_frame = header_box.text_frame
    header_frame.text = "2. 特性规划"
    header_frame.paragraphs[0].font.size = Pt(32)
    header_frame.paragraphs[0].font.bold = True
    header_frame.paragraphs[0].font.color.rgb = PRIMARY_COLOR
    
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(5.5))
    content_frame = content_box.text_frame
    content_frame.text = "本季度立项特性清单（按优先级排序）："
    content_frame.paragraphs[0].font.size = Pt(22)
    content_frame.paragraphs[0].font.bold = True
    
    # 特性表格占位说明
    p = content_frame.add_paragraph()
    p.text = "\n【在此插入特性表格，包含：】"
    p.font.size = Pt(18)
    p.space_before = Pt(15)
    
    items = [
        "• 序号、特性名称、优先级（Must/Should/Could）",
        "• 业务价值、技术复杂度、预计人天",
        "• 负责人、依赖关系、风险等级"
    ]
    for item in items:
        p = content_frame.add_paragraph()
        p.text = item
        p.font.size = Pt(18)
        p.space_before = Pt(6)
    
    # ===== 第5页：资源与负荷 =====
    slide = prs.slides.add_slide(slide_layout)
    
    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    header_frame = header_box.text_frame
    header_frame.text = "3. 资源与负荷分析"
    header_frame.paragraphs[0].font.size = Pt(32)
    header_frame.paragraphs[0].font.bold = True
    header_frame.paragraphs[0].font.color.rgb = PRIMARY_COLOR
    
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(5.5), Inches(5.5))
    content_frame = content_box.text_frame
    content_frame.text = "资源情况："
    content_frame.paragraphs[0].font.size = Pt(22)
    content_frame.paragraphs[0].font.bold = True
    
    resources = [
        "• 季度工作日：【XX】天",
        "• 开发人员系数：【X.X】",
        "• 测试人员系数：【X.X】",
        "• 产品人员系数：【X.X】",
        "• 系数总和：【X.X】"
    ]
    for res in resources:
        p = content_frame.add_paragraph()
        p.text = res
        p.font.size = Pt(18)
        p.space_before = Pt(10)
    
    # 右侧负荷分析
    right_box = slide.shapes.add_textbox(Inches(6.8), Inches(1.5), Inches(5.5), Inches(5.5))
    right_frame = right_box.text_frame
    right_frame.text = "负荷分析："
    right_frame.paragraphs[0].font.size = Pt(22)
    right_frame.paragraphs[0].font.bold = True
    
    loads = [
        "• 总可用产能：【XXX】人天",
        "• 计划人天：【XXX】人天",
        "• 负荷率：【XX%】",
        "• 评估结果：【合理/偏高/偏低】",
        "• 剩余产能：【XXX】人天"
    ]
    for load in loads:
        p = right_frame.add_paragraph()
        p.text = load
        p.font.size = Pt(18)
        p.space_before = Pt(10)
    
    # ===== 第6页：里程碑 =====
    slide = prs.slides.add_slide(slide_layout)
    
    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    header_frame = header_box.text_frame
    header_frame.text = "4. 里程碑与关键节点"
    header_frame.paragraphs[0].font.size = Pt(32)
    header_frame.paragraphs[0].font.bold = True
    header_frame.paragraphs[0].font.color.rgb = PRIMARY_COLOR
    
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(5.5))
    content_frame = content_box.text_frame
    content_frame.text = "关键里程碑："
    content_frame.paragraphs[0].font.size = Pt(22)
    content_frame.paragraphs[0].font.bold = True
    
    milestones = [
        "【月/日】 - 需求评审完成",
        "【月/日】 - 技术方案评审",
        "【月/日】 - 开发完成（转测）",
        "【月/日】 - 测试完成",
        "【月/日】 - 上线发布"
    ]
    for ms in milestones:
        p = content_frame.add_paragraph()
        p.text = f"• {ms}"
        p.font.size = Pt(20)
        p.space_before = Pt(15)
    
    p = content_frame.add_paragraph()
    p.text = "\n【建议在此处添加甘特图或时间线图示】"
    p.font.size = Pt(16)
    p.font.color.rgb = ACCENT_COLOR
    p.space_before = Pt(20)
    
    # ===== 第7页：风险 =====
    slide = prs.slides.add_slide(slide_layout)
    
    header_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(0.8))
    header_frame = header_box.text_frame
    header_frame.text = "5. 风险与应对措施"
    header_frame.paragraphs[0].font.size = Pt(32)
    header_frame.paragraphs[0].font.bold = True
    header_frame.paragraphs[0].font.color.rgb = PRIMARY_COLOR
    
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(5.5))
    content_frame = content_box.text_frame
    content_frame.text = "主要风险："
    content_frame.paragraphs[0].font.size = Pt(22)
    content_frame.paragraphs[0].font.bold = True
    
    risks = [
        ("风险1：【描述风险内容】", "应对：【应对措施】"),
        ("风险2：【描述风险内容】", "应对：【应对措施】"),
        ("风险3：【描述风险内容】", "应对：【应对措施】")
    ]
    
    for risk, solution in risks:
        p = content_frame.add_paragraph()
        p.text = f"• {risk}"
        p.font.size = Pt(18)
        p.space_before = Pt(15)
        
        p = content_frame.add_paragraph()
        p.text = f"  {solution}"
        p.font.size = Pt(16)
        p.font.color.rgb = ACCENT_COLOR
        p.space_before = Pt(4)
    
    # ===== 第8页：结束页 =====
    slide = prs.slides.add_slide(slide_layout)
    
    title_box = slide.shapes.add_textbox(Inches(1), Inches(3), Inches(11.333), Inches(1.5))
    title_frame = title_box.text_frame
    title_frame.text = "谢谢！"
    title_frame.paragraphs[0].font.size = Pt(48)
    title_frame.paragraphs[0].font.bold = True
    title_frame.paragraphs[0].font.color.rgb = PRIMARY_COLOR
    title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(4.5), Inches(11.333), Inches(1))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "Q&A"
    subtitle_frame.paragraphs[0].font.size = Pt(28)
    subtitle_frame.paragraphs[0].font.color.rgb = ACCENT_COLOR
    subtitle_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    
    # 保存
    prs.save(output_path)
    print(f"✅ PPT模板已生成：{output_path}")
    print(f"   包含 8 页幻灯片：")
    print(f"   - 封面")
    print(f"   - 目录")
    print(f"   - 季度目标与策略")
    print(f"   - 特性规划")
    print(f"   - 资源与负荷分析")
    print(f"   - 里程碑与关键节点")
    print(f"   - 风险与应对措施")
    print(f"   - 结束页")


if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="生成季度立项规划报告PPT模板")
    parser.add_argument("--output", "-o", default="planning_template.pptx", help="输出文件路径")
    args = parser.parse_args()
    
    create_planning_template(args.output)
